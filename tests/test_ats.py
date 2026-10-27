"""
tests/test_ats.py : builds every preset, then reads the output back the way a
cheap ATS would (plain text extraction) and checks nothing got mangled.

Run:  pytest -q
      pytest -q -k devops          # one preset
Slow-ish on purpose (LibreOffice runs per fit iteration). Worth it.
"""

from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from build import MASTER, BuildResult, build  # noqa: E402
from models import load_resume  # noqa: E402

RESUME = load_resume(MASTER)
PRESETS = list(RESUME.presets)

# section key -> heading text the template prints (substring match)
HEADINGS = {
    "summary": "PROFESSIONAL SUMMARY",
    "skills": "TECHNICAL SKILLS",
    "experience": "EXPERIENCE",
    "projects": "PROJECTS",
    "education": "EDUCATION",
    "certs": "CERTIFICATIONS",
}

LIGATURES = {"ﬁ": "fi", "ﬂ": "fl", "ﬀ": "ff", "ﬃ": "ffi", "ﬄ": "ffl"}


# ---------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------
def norm(text: str) -> str:
    for k, v in LIGATURES.items():
        text = text.replace(k, v)
    return re.sub(r"\s+", " ", text).strip()


def pdf_text(res: BuildResult) -> str:
    return "\n".join(p.extract_text() or "" for p in PdfReader(str(res.pdf)).pages)


def docx_text(res: BuildResult) -> str:
    return "\n".join(p.text for p in Document(str(res.docx)).paragraphs)


def build_all(out_dir: Path) -> dict[str, BuildResult]:
    return {name: build(RESUME, name, company="ats-test", out_dir=out_dir) for name in PRESETS}


# ---------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------
@pytest.fixture(scope="module")
def built(tmp_path_factory) -> dict[str, BuildResult]:
    return build_all(tmp_path_factory.mktemp("ats"))


# ---------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------
# max_pages is the fit loop's TARGET. Spilling a bit onto page 2 is accepted;
# going past this is not.
HARD_PAGE_CAP = 2


@pytest.mark.parametrize("preset", PRESETS)
def test_fits_page_budget(built, preset):
    res = built[preset]
    assert res.pages <= HARD_PAGE_CAP, (
        f"{preset}: {res.pages} pages (hard cap {HARD_PAGE_CAP}) even after dropping {res.dropped}"
    )


@pytest.mark.parametrize("preset", PRESETS)
def test_no_template_leftovers(built, preset):
    """Unrendered tags, Python reprs, or None leaking into the page."""
    for label, text in (("pdf", pdf_text(built[preset])), ("docx", docx_text(built[preset]))):
        for junk in ("{{", "}}", "{%", "%}", "<built-in", "object at 0x"):
            assert junk not in text, f"{preset} {label}: found '{junk}'"
        assert not re.search(r"\bNone\b", text), f"{preset} {label}: literal 'None' on the page"


@pytest.mark.parametrize("preset", PRESETS)
def test_contact_info_extractable(built, preset):
    text = norm(pdf_text(built[preset]))
    prof = RESUME.profile
    assert prof.name in text
    assert prof.email in text
    digits = re.sub(r"\D", "", prof.phone)
    assert digits in re.sub(r"\D", "", text), "phone number not extractable"
    if prof.links.linkedin:
        assert prof.links.linkedin in text


@pytest.mark.parametrize("preset", PRESETS)
def test_contact_block_is_at_the_top(built, preset):
    text = norm(pdf_text(built[preset]))
    first_heading = min(
        text.find(h) for s, h in HEADINGS.items()
        if s in RESUME.presets[preset].sections and text.find(h) >= 0
    )
    assert text.find(RESUME.profile.email) < first_heading, "contact info must come before sections"


@pytest.mark.parametrize("preset", PRESETS)
def test_sections_in_preset_order(built, preset):
    text = norm(pdf_text(built[preset]))
    p = RESUME.presets[preset]
    expected = [s for s in p.sections if not (s == "projects" and not p.projects)]
    positions = []
    for s in expected:
        pos = text.find(HEADINGS[s])
        assert pos >= 0, f"{preset}: heading for '{s}' missing"
        positions.append(pos)
    assert positions == sorted(positions), f"{preset}: sections out of order {expected}"


@pytest.mark.parametrize("preset", PRESETS)
def test_every_selected_bullet_extractable(built, preset):
    """Catches text boxes, columns, or anything that hides content from parsers."""
    res = built[preset]
    text = norm(pdf_text(res))
    ctx = res.selection.context
    for block in (*ctx["experience"], *ctx["projects"]):
        for b in block["bullets"]:
            snippet = norm(b["text"])[:50]
            assert snippet in text, f"{preset}: bullet '{b['id']}' not extractable"


@pytest.mark.parametrize("preset", PRESETS)
def test_roles_parse_cleanly(built, preset):
    """Title and company must not fuse together ('Technical ConsultantTransact...')."""
    res = built[preset]
    text = norm(pdf_text(res))
    for role in res.selection.context["experience"]:
        assert f"| {role['company']}" in text, f"{preset}: '{role['company']}' fused or missing"
        assert role["dates"].split(" ")[0] in text, f"{preset}: dates missing for {role['company']}"


@pytest.mark.parametrize("preset", PRESETS)
def test_docx_structure_is_ats_safe(built, preset):
    res = built[preset]
    doc = Document(str(res.docx))
    assert not doc.tables, "tables confuse parsers"
    assert not doc.inline_shapes, "no images/QR codes"

    with zipfile.ZipFile(res.docx) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert "w:txbxContent" not in xml, "text boxes hide content from ATS"
    assert "<wp:anchor" not in xml, "floating objects detected"
    cols = re.search(r'<w:cols\b[^>]*\bw:num="(\d+)"', xml)
    assert not cols or int(cols.group(1)) <= 1, "multi-column layout"

    for section in doc.sections:
        for part in (section.header, section.footer):
            assert not any(p.text.strip() for p in part.paragraphs), (
                "content in Word header/footer; some ATS skip it"
            )


@pytest.mark.parametrize("preset", PRESETS)
def test_skills_actually_render(built, preset):
    """Empty skills lines slipped past every other test once. Never again."""
    res = built[preset]
    if "skills" not in RESUME.presets[preset].sections:
        return
    text = norm(pdf_text(res))
    groups = res.selection.context["skills"]
    assert groups, f"{preset}: selector produced zero skill groups"
    for g in groups:
        assert g["names"], f"{preset}: category '{g['category']}' is empty"
        assert norm(g["names"])[:40] in text, (
            f"{preset}: skills for '{g['category']}' missing from the page "
            f"(template/selector key mismatch?)"
        )


@pytest.mark.parametrize("preset", PRESETS)
def test_recruiter_friendly_filename(built, preset):
    res = built[preset]
    assert res.pdf.name.endswith("_Resume.pdf")
    assert " " not in res.pdf.name
    assert RESUME.profile.name.split()[0] in res.pdf.name